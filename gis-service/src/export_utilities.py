"""
Export Utilities
Handles GPX export and polygon vertex export
"""

import logging
import json
from pathlib import Path
from typing import Dict, List
from shapely import wkt
from shapely.geometry import Point, shape
import pyproj
from shapely.ops import transform

logger = logging.getLogger(__name__)


class GPXExporter:
    """Export sample points to GPX format for GPS devices"""
    
    @staticmethod
    def export_sample_points_to_gpx(
        sample_plot_geojson_path: str,
        output_path: str,
        cf_name: str = "Community Forest"
    ) -> str:
        """
        Export sample points to GPX format
        
        Args:
            sample_plot_geojson_path: Path to sample plot GeoJSON
            output_path: Output GPX file path
            cf_name: Community Forest name for metadata
            
        Returns:
            Path to generated GPX file
        """
        logger.info(f"Exporting sample points to GPX: {output_path}")
        
        try:
            # Load sample plot data
            with open(sample_plot_geojson_path, 'r') as f:
                geojson_data = json.load(f)
            
            features = geojson_data.get('features', [])
            
            # Create GPX XML
            gpx_content = GPXExporter._create_gpx_xml(features, cf_name)
            
            # Write to file
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(gpx_content)
            
            logger.info(f"GPX file created with {len(features)} waypoints")
            return output_path
            
        except Exception as e:
            logger.error(f"Error exporting to GPX: {str(e)}")
            raise
    
    @staticmethod
    def _create_gpx_xml(features: List[Dict], cf_name: str) -> str:
        """Create GPX XML content"""
        from datetime import datetime
        
        timestamp = datetime.utcnow().isoformat() + 'Z'
        
        gpx_header = f'''<?xml version="1.0" encoding="UTF-8"?>
<gpx version="1.1" creator="Community Forest Mapping System"
     xmlns="http://www.topografix.com/GPX/1/1"
     xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance"
     xsi:schemaLocation="http://www.topografix.com/GPX/1/1 http://www.topografix.com/GPX/1/1/gpx.xsd">
  <metadata>
    <name>{cf_name} - Sample Points</name>
    <desc>Sample plot locations for forest inventory</desc>
    <time>{timestamp}</time>
  </metadata>
'''
        
        waypoints = []
        for feature in features:
            props = feature.get('properties', {})
            geom = feature.get('geometry', {})
            
            if geom.get('type') != 'Point':
                continue
            
            coords = geom.get('coordinates', [])
            if len(coords) < 2:
                continue
            
            lon, lat = coords[0], coords[1]
            plot_id = props.get('plot_id', 'Unknown')
            compartment = props.get('compartment_label', '')
            subcompartment = props.get('subcompartment_label', '')
            
            waypoint = f'''  <wpt lat="{lat}" lon="{lon}">
    <name>{plot_id}</name>
    <desc>Compartment: {compartment}, Sub-compartment: {subcompartment}</desc>
    <sym>Flag, Blue</sym>
    <type>Sample Plot</type>
  </wpt>
'''
            waypoints.append(waypoint)
        
        gpx_footer = '</gpx>'
        
        return gpx_header + ''.join(waypoints) + gpx_footer


class PolygonVertexExporter:
    """Export polygon vertices to Excel/CSV"""
    
    @staticmethod
    def export_vertices_to_excel(
        compartment_geojson_path: str,
        output_path: str,
        utm_epsg: str = "EPSG:32644"
    ) -> str:
        """
        Export polygon vertices to Excel with UTM coordinates
        
        Args:
            compartment_geojson_path: Path to compartment GeoJSON
            output_path: Output Excel file path
            utm_epsg: UTM projection (default: Zone 44N for Nepal)
            
        Returns:
            Path to generated Excel file
        """
        logger.info(f"Exporting polygon vertices to Excel: {output_path}")
        
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            
            # Load compartment data
            with open(compartment_geojson_path, 'r') as f:
                geojson_data = json.load(f)
            
            features = geojson_data.get('features', [])
            sub_compartments = [f for f in features if f['properties'].get('level') == 1]
            
            # Create transformer for WGS84 to UTM
            transformer = pyproj.Transformer.from_crs(
                "EPSG:4326",  # WGS84
                utm_epsg,
                always_xy=True
            )
            
            # Create workbook
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Polygon Vertices"
            
            # Header row
            headers = ['Compartment ID', 'Sub-compartment ID', 'Vertex Number', 
                      'Easting (m)', 'Northing (m)', 'Coordinate System']
            ws.append(headers)
            
            # Style header
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")
            
            # Extract vertices
            for feature in sub_compartments:
                props = feature['properties']
                compartment_id = props.get('parentLabel', 'C1')
                subcompartment_id = props.get('label', '')
                
                geom = shape(feature['geometry'])
                
                # Get exterior coordinates
                if hasattr(geom, 'exterior'):
                    coords = list(geom.exterior.coords)
                else:
                    continue
                
                # Transform and write vertices
                for idx, (lon, lat) in enumerate(coords[:-1], 1):  # Skip last (duplicate of first)
                    # Transform to UTM
                    easting, northing = transformer.transform(lon, lat)
                    
                    ws.append([
                        compartment_id,
                        subcompartment_id,
                        idx,
                        round(easting, 2),
                        round(northing, 2),
                        utm_epsg
                    ])
            
            # Auto-adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
            
            # Save workbook
            wb.save(output_path)
            logger.info(f"Polygon vertices exported to Excel: {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"Error exporting vertices to Excel: {str(e)}")
            raise
    
    @staticmethod
    def export_vertices_to_csv(
        compartment_geojson_path: str,
        output_path: str,
        utm_epsg: str = "EPSG:32644"
    ) -> str:
        """
        Export polygon vertices to CSV with UTM coordinates
        
        Args:
            compartment_geojson_path: Path to compartment GeoJSON
            output_path: Output CSV file path
            utm_epsg: UTM projection
            
        Returns:
            Path to generated CSV file
        """
        logger.info(f"Exporting polygon vertices to CSV: {output_path}")
        
        try:
            import csv
            
            # Load compartment data
            with open(compartment_geojson_path, 'r') as f:
                geojson_data = json.load(f)
            
            features = geojson_data.get('features', [])
            sub_compartments = [f for f in features if f['properties'].get('level') == 1]
            
            # Create transformer
            transformer = pyproj.Transformer.from_crs(
                "EPSG:4326",
                utm_epsg,
                always_xy=True
            )
            
            # Write CSV
            with open(output_path, 'w', newline='', encoding='utf-8') as csvfile:
                writer = csv.writer(csvfile)
                
                # Header
                writer.writerow(['Compartment ID', 'Sub-compartment ID', 'Vertex Number',
                               'Easting (m)', 'Northing (m)', 'Coordinate System'])
                
                # Extract vertices
                for feature in sub_compartments:
                    props = feature['properties']
                    compartment_id = props.get('parentLabel', 'C1')
                    subcompartment_id = props.get('label', '')
                    
                    geom = shape(feature['geometry'])
                    
                    if hasattr(geom, 'exterior'):
                        coords = list(geom.exterior.coords)
                    else:
                        continue
                    
                    for idx, (lon, lat) in enumerate(coords[:-1], 1):
                        easting, northing = transformer.transform(lon, lat)
                        
                        writer.writerow([
                            compartment_id,
                            subcompartment_id,
                            idx,
                            round(easting, 2),
                            round(northing, 2),
                            utm_epsg
                        ])
            
            logger.info(f"Polygon vertices exported to CSV: {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"Error exporting vertices to CSV: {str(e)}")
            raise


class EnhancedSamplePlotExporter:
    """Enhanced sample plot export with sub-compartment information"""
    
    @staticmethod
    def export_to_excel(
        sample_plot_geojson_path: str,
        output_path: str,
        utm_epsg: str = "EPSG:32644"
    ) -> str:
        """
        Export sample plots to Excel with enhanced format
        
        Args:
            sample_plot_geojson_path: Path to sample plot GeoJSON
            output_path: Output Excel file path
            utm_epsg: UTM projection
            
        Returns:
            Path to generated Excel file
        """
        logger.info(f"Exporting sample plots to Excel: {output_path}")
        
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            
            # Load sample plot data
            with open(sample_plot_geojson_path, 'r') as f:
                geojson_data = json.load(f)
            
            features = geojson_data.get('features', [])
            
            # Create transformer
            transformer = pyproj.Transformer.from_crs(
                "EPSG:4326",
                utm_epsg,
                always_xy=True
            )
            
            # Create workbook
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sample Plots"
            
            # Header row
            headers = ['Sample Point ID', 'Compartment ID', 'Sub-compartment ID',
                      'Latitude (WGS84)', 'Longitude (WGS84)', 
                      'Easting (m)', 'Northing (m)', 'Coordinate System']
            ws.append(headers)
            
            # Style header
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="2196F3", end_color="2196F3", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")
            
            # Write data
            for feature in features:
                props = feature.get('properties', {})
                geom = feature.get('geometry', {})
                
                if geom.get('type') != 'Point':
                    continue
                
                coords = geom.get('coordinates', [])
                if len(coords) < 2:
                    continue
                
                lon, lat = coords[0], coords[1]
                
                # Transform to UTM
                easting, northing = transformer.transform(lon, lat)
                
                ws.append([
                    props.get('plot_id', ''),
                    props.get('compartment_label', ''),
                    props.get('subcompartment_label', ''),
                    round(lat, 6),
                    round(lon, 6),
                    round(easting, 2),
                    round(northing, 2),
                    utm_epsg
                ])
            
            # Auto-adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
            
            # Save workbook
            wb.save(output_path)
            logger.info(f"Sample plots exported to Excel: {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"Error exporting sample plots to Excel: {str(e)}")
            raise
